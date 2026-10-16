#!/usr/bin/env python3
"""
TikTok comment watcher - 24/7 background worker (Render.com)

MODE=followers (default when FOLLOWERS_FILE exists)
    Read a list of usernames (xlsx / csv / txt), open each public profile,
    list its videos, open every video that has comments, capture the comment
    API responses and look for comments written by TARGET_USERNAME.

MODE=discover
    Discover videos through search queries / explore page and scan them.

Matches are appended to Google Sheets; SQLite keeps history so that nothing is
scanned or saved twice and the job can resume after every restart.

Use only on public data and in line with TikTok's Terms of Service / local law.
"""

import asyncio
import base64
import csv
import hashlib
import json
import logging
import os
import random
import re
import signal
import sqlite3
import sys
import time
from contextlib import closing
from datetime import datetime, timezone
from typing import Dict, List, Optional, Tuple

import gspread
from google.oauth2.service_account import Credentials
from playwright.async_api import Browser, BrowserContext, Page, async_playwright

try:
    from fake_useragent import UserAgent
except Exception:  # pragma: no cover
    UserAgent = None


# --------------------------------------------------------------------------- #
# Configuration (environment variables)
# --------------------------------------------------------------------------- #
def _env_int(name: str, default: int) -> int:
    try:
        return int(os.getenv(name, default))
    except ValueError:
        return default


TARGET_USERNAME = (os.getenv("TARGET_USERNAME") or "i7si_u").lstrip("@").strip().lower()  # empty secret -> default
GOOGLE_SHEET_NAME = os.getenv("GOOGLE_SHEET_NAME", "")
GOOGLE_SHEET_ID = os.getenv("GOOGLE_SHEET_ID", "")
GOOGLE_WORKSHEET = os.getenv("GOOGLE_WORKSHEET", "")
GOOGLE_CREDENTIALS_JSON = os.getenv("GOOGLE_CREDENTIALS_JSON", "")  # raw JSON or base64
GOOGLE_CREDENTIALS_FILE = os.getenv("GOOGLE_CREDENTIALS_FILE", "google_credentials.json")  # local fallback
HEADLESS = os.getenv("HEADLESS", "true").lower() != "false"  # "false" + xvfb-run = less detectable
WARMUP = os.getenv("WARMUP", "true").lower() == "true"  # visit tiktok.com first to get normal cookies
PROFILE_RELOADS = _env_int("PROFILE_RELOADS", 2)  # page reloads when TikTok says "Something went wrong"
BLOCK_LIMIT = _env_int("BLOCK_LIMIT", 8)  # consecutive blocked profiles before the run aborts
DEBUG_DIR = os.getenv("DEBUG_DIR", "debug")
DEBUG_DUMPS_MAX = _env_int("DEBUG_DUMPS_MAX", 3)  # screenshots/log dumps of blocked pages
RUN_ONCE = os.getenv("RUN_ONCE", "false").lower() == "true"  # GitHub Actions: one pass then exit
MAX_RUNTIME_MINUTES = _env_int("MAX_RUNTIME_MINUTES", 0)  # 0 = unlimited (Render worker)

FOLLOWERS_FILE = os.getenv("FOLLOWERS_FILE", "followers.csv")
MODE = (os.getenv("MODE") or ("followers" if os.path.exists(FOLLOWERS_FILE) else "discover")).lower()

# discover mode
SEARCH_QUERIES = [q.strip() for q in os.getenv("SEARCH_QUERIES", TARGET_USERNAME).split(",") if q.strip()]
USE_EXPLORE = os.getenv("USE_EXPLORE", "true").lower() == "true"
DISCOVERY_SCROLLS = _env_int("DISCOVERY_SCROLLS", 6)
MAX_VIDEOS_PER_CYCLE = _env_int("MAX_VIDEOS_PER_CYCLE", 25)

# followers mode
CONCURRENCY = max(1, _env_int("CONCURRENCY", 2))  # parallel browser contexts
PROFILE_MAX_VIDEOS = _env_int("PROFILE_MAX_VIDEOS", 40)  # videos listed per profile (newest first)
MAX_VIDEOS_PER_PROFILE = _env_int("MAX_VIDEOS_PER_PROFILE", 20)  # videos actually opened per profile
PROFILE_SCROLLS = _env_int("PROFILE_SCROLLS", 8)
PROFILE_RESCAN_HOURS = _env_int("PROFILE_RESCAN_HOURS", 168)  # re-check finished profiles weekly
MAX_ATTEMPTS = _env_int("MAX_ATTEMPTS", 3)
RECYCLE_CONTEXT_EVERY = _env_int("RECYCLE_CONTEXT_EVERY", 20)  # profiles per fresh context

# common
COMMENT_SCROLLS = _env_int("COMMENT_SCROLLS", 15)
DB_PATH = os.getenv("DB_PATH", "history.db")
CYCLE_SLEEP_SECONDS = _env_int("CYCLE_SLEEP_SECONDS", 900)
RESCAN_AFTER_HOURS = _env_int("RESCAN_AFTER_HOURS", 72)
PURGE_AFTER_DAYS = _env_int("PURGE_AFTER_DAYS", 60)
RESTART_BROWSER_EVERY = _env_int("RESTART_BROWSER_EVERY", 5)  # cycles
BACKOFF_BASE = _env_int("BACKOFF_BASE_SECONDS", 30)
BACKOFF_MAX = _env_int("BACKOFF_MAX_SECONDS", 1800)
NAV_TIMEOUT_MS = _env_int("NAV_TIMEOUT_MS", 45000)
PROXY_SERVER = os.getenv("PROXY_SERVER", "")
PROXY_USERNAME = os.getenv("PROXY_USERNAME", "")
PROXY_PASSWORD = os.getenv("PROXY_PASSWORD", "")

SHEET_HEADER = ["Timestamp", "Video URL", "Target Username", "Comment Text", "Likes", "Video Owner"]
VIDEO_RE = re.compile(r"^https://www\.tiktok\.com/@[\w.\-]+/video/\d+")

logging.basicConfig(
    level=os.getenv("LOG_LEVEL", "INFO"),
    format="%(asctime)s | %(levelname)-7s | %(message)s",
    stream=sys.stdout,
)
log = logging.getLogger("tiktok-worker")
logging.getLogger("fake_useragent").setLevel(logging.ERROR)  # silence harmless fallback warnings
STOP = asyncio.Event()
BLOCK_STATE = {"streak": 0, "tripped": False, "dumps": 0}


class BlockedError(Exception):
    """Captcha / soft-block / empty response where content was expected."""


# --------------------------------------------------------------------------- #
# SQLite history
# --------------------------------------------------------------------------- #
class History:
    def __init__(self, path: str):
        self.path = path
        with closing(self._conn()) as c:
            c.executescript(
                """
                CREATE TABLE IF NOT EXISTS videos (
                    url TEXT PRIMARY KEY,
                    scanned_at INTEGER NOT NULL,
                    comment_count INTEGER NOT NULL DEFAULT -1
                );
                CREATE TABLE IF NOT EXISTS comments (
                    cid TEXT PRIMARY KEY,
                    video_url TEXT NOT NULL,
                    saved_at INTEGER NOT NULL
                );
                CREATE TABLE IF NOT EXISTS profiles (
                    username TEXT PRIMARY KEY,
                    status TEXT NOT NULL,
                    videos_found INTEGER NOT NULL DEFAULT 0,
                    scanned_at INTEGER NOT NULL
                );
                """
            )
            try:  # migrate databases created by the older version
                c.execute("ALTER TABLE videos ADD COLUMN comment_count INTEGER NOT NULL DEFAULT -1")
            except sqlite3.OperationalError:
                pass
            c.commit()

    def _conn(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.path, timeout=30)
        conn.execute("PRAGMA journal_mode=WAL")
        return conn

    # ---- videos
    def video_recently_scanned(self, url: str, within_hours: int) -> bool:
        cutoff = int(time.time()) - within_hours * 3600
        with closing(self._conn()) as c:
            row = c.execute("SELECT scanned_at FROM videos WHERE url=?", (url,)).fetchone()
        return bool(row and row[0] >= cutoff)

    def video_needs_scan(self, url: str, comment_count: int) -> bool:
        """Scan if never scanned, or if the comment count grew since last scan."""
        with closing(self._conn()) as c:
            row = c.execute("SELECT scanned_at, comment_count FROM videos WHERE url=?", (url,)).fetchone()
        if row is None:
            return True
        scanned_at, stored = row
        if comment_count >= 0 and stored >= 0:
            return comment_count > stored
        return scanned_at < int(time.time()) - RESCAN_AFTER_HOURS * 3600

    def mark_video(self, url: str, comment_count: int = -1) -> None:
        with closing(self._conn()) as c:
            c.execute(
                "INSERT OR REPLACE INTO videos(url, scanned_at, comment_count) VALUES(?,?,?)",
                (url, int(time.time()), comment_count),
            )
            c.commit()

    # ---- comments
    def comment_seen(self, cid: str) -> bool:
        with closing(self._conn()) as c:
            return c.execute("SELECT 1 FROM comments WHERE cid=?", (cid,)).fetchone() is not None

    def mark_comment(self, cid: str, video_url: str) -> None:
        with closing(self._conn()) as c:
            c.execute(
                "INSERT OR IGNORE INTO comments(cid, video_url, saved_at) VALUES(?,?,?)",
                (cid, video_url, int(time.time())),
            )
            c.commit()

    # ---- profiles
    def profile_due(self, username: str) -> bool:
        with closing(self._conn()) as c:
            row = c.execute("SELECT status, scanned_at FROM profiles WHERE username=?", (username,)).fetchone()
        if row is None:
            return True
        status, scanned_at = row
        if status == "error":
            return True
        return scanned_at < int(time.time()) - PROFILE_RESCAN_HOURS * 3600

    def set_profile(self, username: str, status: str, videos_found: int = 0) -> None:
        with closing(self._conn()) as c:
            c.execute(
                "INSERT OR REPLACE INTO profiles(username, status, videos_found, scanned_at) VALUES(?,?,?,?)",
                (username, status, videos_found, int(time.time())),
            )
            c.commit()

    def purge(self, days: int) -> None:
        """Periodic cache clearing of old history (comment ids are kept the longest)."""
        cutoff = int(time.time()) - days * 86400
        with closing(self._conn()) as c:
            v = c.execute("DELETE FROM videos WHERE scanned_at < ?", (cutoff,)).rowcount
            k = c.execute("DELETE FROM comments WHERE saved_at < ?", (cutoff - 30 * 86400,)).rowcount
            c.commit()
            c.execute("VACUUM")
        log.info("History purge: removed %s videos, %s comments", v, k)


# --------------------------------------------------------------------------- #
# Google Sheets
# --------------------------------------------------------------------------- #
class Sheet:
    SCOPES = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]

    def __init__(self):
        self.ws: Optional[gspread.Worksheet] = None
        self.keys: set = set()

    @staticmethod
    def make_key(video_url: str, text: str) -> str:
        return hashlib.sha1(f"{video_url}|{text.strip()}".encode()).hexdigest()

    def _load_keys(self) -> None:
        """Remember what is already in the sheet, so duplicates are impossible even without history.db."""
        try:
            for r in self.ws.get_all_values()[1:]:
                if len(r) >= 4:
                    self.keys.add(self.make_key(r[1], r[3]))
        except Exception as e:  # not fatal
            log.warning("Could not preload existing sheet rows: %s", e)

    @staticmethod
    def _load_credentials() -> dict:
        raw = GOOGLE_CREDENTIALS_JSON.strip()
        if not raw and os.path.exists(GOOGLE_CREDENTIALS_FILE):
            with open(GOOGLE_CREDENTIALS_FILE, encoding="utf-8") as f:
                raw = f.read().strip()
        if not raw:
            raise RuntimeError("Set GOOGLE_CREDENTIALS_JSON (or provide google_credentials.json)")
        if not raw.startswith("{"):
            raw = base64.b64decode(raw).decode("utf-8")
        return json.loads(raw)

    def connect(self) -> None:
        creds = Credentials.from_service_account_info(self._load_credentials(), scopes=self.SCOPES)
        gc = gspread.authorize(creds)
        if GOOGLE_SHEET_ID:
            sh = gc.open_by_key(GOOGLE_SHEET_ID)
        elif GOOGLE_SHEET_NAME:
            sh = gc.open(GOOGLE_SHEET_NAME)
        else:
            raise RuntimeError("Set GOOGLE_SHEET_NAME or GOOGLE_SHEET_ID")
        self.ws = sh.worksheet(GOOGLE_WORKSHEET) if GOOGLE_WORKSHEET else sh.sheet1
        first = self.ws.row_values(1)
        if not first:
            self.ws.append_row(SHEET_HEADER, value_input_option="RAW")
        elif first == SHEET_HEADER[:5]:  # old 5-column header -> add the 6th column
            self.ws.update(values=[SHEET_HEADER], range_name="A1:F1")
        elif first != SHEET_HEADER:
            self.ws.insert_row(SHEET_HEADER, 1, value_input_option="RAW")
        self._load_keys()
        log.info("Connected to Google Sheet: %s (%d existing rows indexed)", sh.title, len(self.keys))

    def append(self, row: List) -> None:
        if self.ws is None:
            self.connect()
        try:
            self.ws.append_row(row, value_input_option="RAW", table_range="A1")  # RAW: no formula injection
        except gspread.exceptions.APIError:
            self.connect()
            self.ws.append_row(row, value_input_option="RAW", table_range="A1")
        self.keys.add(self.make_key(str(row[1]), str(row[3])))


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def random_user_agent(real_version: str = "") -> str:
    """fake-useragent pick, but only if its Chrome version is close to the real browser (consistency)."""
    consistent = (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
        f"Chrome/{real_version or '129.0.0.0'} Safari/537.36"
    )
    if UserAgent is None:
        return consistent
    try:
        ua = UserAgent(browsers=["Chrome"], os=["Windows", "Mac OS X"], platforms=["desktop"]).random
        m = re.search(r"Chrome/(\d+)", ua)
        real_major = int(real_version.split(".")[0]) if real_version else 0
        if m and "Edg" not in ua and (not real_major or abs(int(m.group(1)) - real_major) <= 6):
            return ua
    except Exception:
        pass
    return consistent


def clean_video_url(href: str) -> Optional[str]:
    if href.startswith("/"):
        href = "https://www.tiktok.com" + href
    href = href.split("?")[0]
    m = VIDEO_RE.match(href)
    return m.group(0) if m else None


def parse_count(text: str) -> int:
    t = (text or "").strip().upper().replace(",", "")
    if not t:
        return 0
    mult = 1
    if t.endswith("K"):
        mult, t = 1_000, t[:-1]
    elif t.endswith("M"):
        mult, t = 1_000_000, t[:-1]
    elif t.endswith("B"):
        mult, t = 1_000_000_000, t[:-1]
    try:
        return int(float(t) * mult)
    except ValueError:
        return 0


async def human_pause(lo: float = 0.8, hi: float = 2.2) -> None:
    await asyncio.sleep(random.uniform(lo, hi))


async def interruptible_sleep(seconds: float) -> None:
    try:
        await asyncio.wait_for(STOP.wait(), timeout=seconds)
    except asyncio.TimeoutError:
        pass


def load_followers(path: str) -> List[dict]:
    """Load usernames from .csv / .xlsx / .txt. Returns [{'username','videos','private'}]."""
    name_keys = ("authorMeta/name", "username", "unique_id", "uniqueId", "name", "user")
    video_keys = ("authorMeta/video", "videos", "video_count")
    private_keys = ("authorMeta/privateAccount", "private", "privateAccount")

    def truthy(v) -> bool:
        return str(v).strip().lower() in {"true", "1", "yes"}

    def pick(row: dict, keys):
        for k in keys:
            if k in row and row[k] not in (None, ""):
                return row[k]
        return None

    rows: List[dict] = []
    ext = os.path.splitext(path)[1].lower()
    if ext == ".txt":
        with open(path, encoding="utf-8") as f:
            raw = [{"username": line.strip()} for line in f if line.strip()]
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        it = wb.active.iter_rows(values_only=True)
        header = [str(h) if h is not None else "" for h in next(it)]
        raw = [dict(zip(header, r)) for r in it]
        wb.close()
    else:
        with open(path, encoding="utf-8-sig", newline="") as f:
            raw = list(csv.DictReader(f))

    seen = set()
    for r in raw:
        u = pick(r, name_keys)
        if not u:
            continue
        u = str(u).strip().lstrip("@")
        if not u or u.lower() in seen:
            continue
        seen.add(u.lower())
        vids = pick(r, video_keys)
        try:
            vids = int(float(vids)) if vids is not None else -1
        except ValueError:
            vids = -1
        rows.append({"username": u, "videos": vids, "private": truthy(pick(r, private_keys) or False)})
    return rows


# --------------------------------------------------------------------------- #
# Browser
# --------------------------------------------------------------------------- #
async def launch_browser(pw) -> Browser:
    kwargs = dict(
        headless=HEADLESS,
        args=[
            "--no-sandbox",
            "--disable-dev-shm-usage",
            "--disable-gpu",
            "--disable-blink-features=AutomationControlled",
            "--mute-audio",
        ],
    )
    if PROXY_SERVER:
        kwargs["proxy"] = {
            "server": PROXY_SERVER,
            "username": PROXY_USERNAME or None,
            "password": PROXY_PASSWORD or None,
        }
    return await pw.chromium.launch(**kwargs)


async def new_context(browser: Browser) -> BrowserContext:
    ctx = await browser.new_context(
        user_agent=random_user_agent(browser.version),
        viewport={"width": random.choice([1280, 1366, 1440, 1536]), "height": random.choice([720, 800, 900])},
        locale="en-US",
        timezone_id="UTC",
    )
    ctx.set_default_navigation_timeout(NAV_TIMEOUT_MS)
    ctx.set_default_timeout(15000)

    async def _route(route):  # save memory / bandwidth
        if route.request.resource_type in {"image", "media", "font"}:
            await route.abort()
        else:
            await route.continue_()

    await ctx.route("**/*", _route)
    await ctx.add_init_script("Object.defineProperty(navigator,'webdriver',{get:()=>undefined});")
    return ctx


async def warm_up(ctx: BrowserContext) -> None:
    """Open the home page once so the context gets the normal TikTok cookies/tokens before scraping."""
    if not WARMUP:
        return
    page = await ctx.new_page()
    try:
        await page.goto("https://www.tiktok.com/", wait_until="domcontentloaded")
        await human_pause(3, 5)
        await dismiss_popups(page)
        await page.mouse.wheel(0, random.randint(400, 900))
        await human_pause(1.5, 3)
    except Exception as e:
        log.debug("warm-up failed: %s", e)
    finally:
        await page.close()


async def dismiss_popups(page: Page) -> None:
    for sel in (
        'button:has-text("Accept all")',
        'button:has-text("Decline optional cookies")',
        '[data-e2e="modal-close-inner-button"]',
        'div[role="dialog"] button[aria-label="Close"]',
    ):
        try:
            loc = page.locator(sel).first
            if await loc.is_visible(timeout=800):
                await loc.click(timeout=1500)
        except Exception:
            pass


async def page_text(page: Page) -> str:
    try:
        return (await page.inner_text("body", timeout=3000))[:3000].lower()
    except Exception:
        return ""


async def looks_blocked(page: Page) -> bool:
    try:
        if await page.locator("#captcha_container, .captcha_verify_container, [id*='captcha']").count():
            return True
    except Exception:
        pass
    text = await page_text(page)
    return "verify to continue" in text or ("something went wrong" in text and "refresh" in text)


async def dump_debug(page: Page, tag: str, extra: str = "") -> None:
    """Log what TikTok really returned (captcha? login wall? empty?) + screenshot, a few times per run."""
    if BLOCK_STATE["dumps"] >= DEBUG_DUMPS_MAX:
        return
    BLOCK_STATE["dumps"] += 1
    try:
        title = await page.title()
    except Exception:
        title = "?"
    body = " ".join((await page_text(page))[:400].split())
    log.warning("DEBUG %s | url=%s | title=%r | body=%r | api_calls=%s", tag, page.url, title, body, extra)
    try:
        os.makedirs(DEBUG_DIR, exist_ok=True)
        await page.screenshot(path=f"{DEBUG_DIR}/{tag}.png")
    except Exception:
        pass


# --------------------------------------------------------------------------- #
# Comment extraction (one video)
# --------------------------------------------------------------------------- #
async def extract_dom_comments(page: Page, video_url: str) -> List[dict]:
    """DOM fallback: comment containers that include a link to TARGET_USERNAME."""
    results = []
    user_links = page.locator(f'a[href="/@{TARGET_USERNAME}"], a[href^="/@{TARGET_USERNAME}?"]')
    for i in range(await user_links.count()):
        try:
            container = user_links.nth(i).locator(
                'xpath=ancestor::div[contains(@class,"CommentItem") or contains(@class,"DivCommentItemContainer") '
                'or contains(@data-e2e,"comment")][1]'
            )
            if await container.count() == 0:
                continue
            text_loc = container.locator(
                '[data-e2e="comment-level-1"], [data-e2e="comment-level-2"], span[data-e2e^="comment-level"]'
            ).first
            text = (await text_loc.inner_text(timeout=1500)).strip() if await text_loc.count() else ""
            likes_loc = container.locator('[data-e2e="comment-like-count"], span[class*="LikeCount"]').first
            likes = parse_count(await likes_loc.inner_text(timeout=1000)) if await likes_loc.count() else 0
            if text:
                cid = "dom_" + hashlib.sha1(f"{video_url}|{text}".encode()).hexdigest()[:20]
                results.append({"cid": cid, "text": text, "likes": likes})
        except Exception:
            continue
    return results


async def scan_video(ctx: BrowserContext, video_url: str) -> List[dict]:
    """Return the comments written by TARGET_USERNAME on one video."""
    page = await ctx.new_page()
    matches: Dict[str, dict] = {}
    state = {"pages": 0, "has_more": True}
    pending: List[asyncio.Task] = []

    async def on_response(resp):
        try:
            if "/api/comment/list" not in resp.url or resp.status != 200:
                return
            data = await resp.json()
            state["pages"] += 1
            state["has_more"] = bool(data.get("has_more"))
            for c in data.get("comments") or []:
                uname = ((c.get("user") or {}).get("unique_id") or "").lower()
                if uname == TARGET_USERNAME:
                    cid = str(c.get("cid"))
                    matches[cid] = {"cid": cid, "text": (c.get("text") or "").strip(), "likes": int(c.get("digg_count") or 0)}
                # replies embedded in the comment object
                for rc in c.get("reply_comment") or []:
                    if ((rc.get("user") or {}).get("unique_id") or "").lower() == TARGET_USERNAME:
                        rid = str(rc.get("cid"))
                        matches[rid] = {"cid": rid, "text": (rc.get("text") or "").strip(), "likes": int(rc.get("digg_count") or 0)}
        except Exception:
            pass

    page.on("response", lambda r: pending.append(asyncio.create_task(on_response(r))))

    try:
        await page.goto(video_url, wait_until="domcontentloaded")
        await human_pause(2, 3.5)
        await dismiss_popups(page)

        try:  # open the comment panel if collapsed
            icon = page.locator('[data-e2e="comment-icon"]').first
            if await icon.is_visible(timeout=2500):
                await icon.click(timeout=2000)
                await human_pause(1, 1.5)
        except Exception:
            pass

        container = page.locator('[data-e2e="comment-list"], div[class*="DivCommentListContainer"]').first
        idle = 0
        for _ in range(COMMENT_SCROLLS):
            if STOP.is_set():
                break
            before = state["pages"]
            try:
                if await container.count():
                    await container.evaluate("el => el.scrollBy(0, el.clientHeight * 0.9)")
                else:
                    await page.mouse.wheel(0, 1200)
            except Exception:
                await page.mouse.wheel(0, 1200)
            await human_pause(0.9, 1.6)
            if state["pages"] > 0 and not state["has_more"]:
                break  # reached the last page of comments
            idle = idle + 1 if state["pages"] == before else 0
            if idle >= 3:
                break

        if pending:
            await asyncio.gather(*pending, return_exceptions=True)

        if state["pages"] == 0 and await looks_blocked(page):
            raise BlockedError(f"captcha/block on {video_url}")

        if not matches and state["pages"] == 0:  # API never fired -> DOM fallback
            for c in await extract_dom_comments(page, video_url):
                matches[c["cid"]] = c
        return list(matches.values())
    finally:
        await page.close()


async def save_matches(sheet: Sheet, sheet_lock: asyncio.Lock, history: History,
                       comments: List[dict], video_url: str, owner: str) -> int:
    saved = 0
    for c in comments:
        if history.comment_seen(c["cid"]):
            continue
        if Sheet.make_key(video_url, c["text"]) in sheet.keys:  # already in the sheet
            history.mark_comment(c["cid"], video_url)
            continue
        row = [
            datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
            video_url,
            f"@{TARGET_USERNAME}",
            c["text"],
            c["likes"],
            f"@{owner}" if owner else "",
        ]
        async with sheet_lock:
            await asyncio.to_thread(sheet.append, row)  # raises on failure -> not marked as seen
        history.mark_comment(c["cid"], video_url)
        saved += 1
        log.info("MATCH saved: %s commented on %s", TARGET_USERNAME, video_url)
    return saved


# --------------------------------------------------------------------------- #
# MODE = followers
# --------------------------------------------------------------------------- #
async def scan_profile_videos(ctx: BrowserContext, username: str, expected: int) -> Tuple[List[dict], str]:
    """Return ([{'url','comments'}], status) with status in ok / private / empty."""
    page = await ctx.new_page()
    items: Dict[str, dict] = {}
    state = {"responses": 0, "has_more": True}
    pending: List[asyncio.Task] = []
    diag: List[str] = []  # what the post-list API really answered (for DEBUG)

    async def on_response(resp):
        try:
            if "/api/post/item_list" not in resp.url:
                return
            body = await resp.text()
            if resp.status != 200 or not body.strip():
                if len(diag) < 6:
                    diag.append(f"http{resp.status}:len{len(body)}")
                return
            data = json.loads(body)
            n = len(data.get("itemList") or [])
            if len(diag) < 6:
                diag.append(f"ok:code={data.get('statusCode')},msg={str(data.get('status_msg') or '')[:25]!r},items={n},more={data.get('hasMore')}")
            state["responses"] += 1
            state["has_more"] = bool(data.get("hasMore"))
            for it in data.get("itemList") or []:
                vid = str(it.get("id"))
                cnt = int((it.get("stats") or {}).get("commentCount") or 0)
                items[vid] = {"url": f"https://www.tiktok.com/@{username}/video/{vid}", "comments": cnt}
        except Exception as e:
            if len(diag) < 6:
                diag.append(f"err:{type(e).__name__}")

    page.on("response", lambda r: pending.append(asyncio.create_task(on_response(r))))
    dom_links: Dict[str, None] = {}

    async def collect() -> None:
        stagnant, last_total = 0, -1
        for _ in range(PROFILE_SCROLLS):
            hrefs = await page.eval_on_selector_all(
                'a[href*="/video/"]', "els => els.map(e => e.getAttribute('href'))"
            )
            for h in hrefs:
                u = clean_video_url(h or "")
                if u:
                    dom_links[u] = None
            total = len(set(i["url"] for i in items.values()) | set(dom_links))
            if total >= PROFILE_MAX_VIDEOS:
                break
            stagnant = stagnant + 1 if total == last_total else 0
            last_total = total
            if stagnant >= 2 or (state["responses"] > 0 and not state["has_more"] and total > 0):
                break
            await page.mouse.wheel(0, random.randint(1800, 2800))
            await human_pause(1.0, 2.0)

    try:
        await page.goto(f"https://www.tiktok.com/@{username}", wait_until="domcontentloaded")
        await human_pause(2, 3.5)
        await dismiss_popups(page)

        text = await page_text(page)
        if "this account is private" in text:
            return [], "private"
        if "couldn't find this account" in text:
            return [], "empty"

        await collect()

        # TikTok sometimes answers "Something went wrong ... refresh": try reloading like a human would
        for attempt in range(PROFILE_RELOADS):
            if items or dom_links:
                break
            if pending:
                await asyncio.gather(*pending, return_exceptions=True)
            await human_pause(3, 6)
            try:
                btn = page.locator('button:has-text("Refresh")').first
                if await btn.is_visible(timeout=1000):
                    await btn.click(timeout=2000)
                else:
                    await page.reload(wait_until="domcontentloaded")
            except Exception:
                await page.reload(wait_until="domcontentloaded")
            await human_pause(3, 5)
            await collect()

        if pending:
            await asyncio.gather(*pending, return_exceptions=True)

        videos = list(items.values())
        known = {v["url"] for v in videos}
        videos += [{"url": u, "comments": -1} for u in dom_links if u not in known]  # comment count unknown

        if not videos:
            if await looks_blocked(page) or expected > 0:
                await dump_debug(page, f"profile_{username}", "item_list=[" + "; ".join(diag) + "]" if diag else "item_list=none")
                raise BlockedError(f"no videos returned for @{username} (expected {expected})")
            return [], "empty"
        return videos[:PROFILE_MAX_VIDEOS], "ok"
    finally:
        await page.close()


async def process_profile(ctx, username, expected, history, sheet, sheet_lock, stats) -> None:
    videos, status = await scan_profile_videos(ctx, username, expected)
    if status in ("private", "empty"):
        history.set_profile(username, status, 0)
        return

    todo = [v for v in videos if v["comments"] != 0 and history.video_needs_scan(v["url"], v["comments"])]
    for v in todo[:MAX_VIDEOS_PER_PROFILE]:
        if STOP.is_set():
            return  # profile stays "not done" and is resumed after restart
        comments = await scan_video(ctx, v["url"])
        stats["matches"] += await save_matches(sheet, sheet_lock, history, comments, v["url"], username)
        history.mark_video(v["url"], v["comments"])
        stats["videos"] += 1
        await human_pause(1.5, 3.5)
    history.set_profile(username, "done", len(videos))


async def followers_worker(wid: int, queue: asyncio.Queue, browser: Browser, history: History,
                           sheet: Sheet, sheet_lock: asyncio.Lock, stats: dict) -> None:
    ctx = await new_context(browser)
    await warm_up(ctx)
    processed, fails = 0, 0
    try:
        while not STOP.is_set():
            try:
                username, expected, attempt = queue.get_nowait()
            except asyncio.QueueEmpty:
                break
            try:
                await process_profile(ctx, username, expected, history, sheet, sheet_lock, stats)
                fails = 0
                BLOCK_STATE["streak"] = 0
                stats["profiles"] += 1
                log.info("[w%d] @%s done | profiles %d/%d | videos %d | matches %d",
                         wid, username, stats["profiles"], stats["total"], stats["videos"], stats["matches"])
            except Exception as e:
                fails += 1
                if isinstance(e, BlockedError):
                    BLOCK_STATE["streak"] += 1
                    if BLOCK_STATE["streak"] >= BLOCK_LIMIT and not BLOCK_STATE["tripped"]:
                        BLOCK_STATE["tripped"] = True
                        log.error("TikTok is blocking this IP (%d blocked profiles in a row). Stopping the run. "
                                  "Use a residential proxy (PROXY_SERVER) and see the DEBUG lines above.", BLOCK_STATE["streak"])
                        STOP.set()
                if attempt + 1 < MAX_ATTEMPTS:
                    await queue.put((username, expected, attempt + 1))
                else:
                    history.set_profile(username, "error", 0)
                    stats["profiles"] += 1
                delay = min(BACKOFF_BASE * (2 ** (fails - 1)), BACKOFF_MAX) + random.uniform(0, 10)
                log.warning("[w%d] @%s failed (%s: %s) -> backoff %.0fs", wid, username, type(e).__name__, e, delay)
                try:
                    await ctx.close()
                except Exception:
                    pass
                await interruptible_sleep(delay)
                ctx = await new_context(browser)
                await warm_up(ctx)
            processed += 1
            if processed % RECYCLE_CONTEXT_EVERY == 0:  # periodic cache clearing
                await ctx.close()
                ctx = await new_context(browser)
                await warm_up(ctx)
    finally:
        try:
            await ctx.close()
        except Exception:
            pass


async def run_followers_cycle(browser: Browser, history: History, sheet: Sheet) -> int:
    people = load_followers(FOLLOWERS_FILE)
    eligible = [p for p in people if not p["private"] and p["videos"] != 0]
    due = [p for p in eligible if history.profile_due(p["username"])]
    log.info("Followers file: %d total | %d public with videos | %d due this cycle",
             len(people), len(eligible), len(due))
    if not due:
        return 0

    queue: asyncio.Queue = asyncio.Queue()
    for p in due:
        queue.put_nowait((p["username"], p["videos"], 0))
    stats = {"profiles": 0, "videos": 0, "matches": 0, "total": len(due)}
    sheet_lock = asyncio.Lock()

    workers = [
        asyncio.create_task(followers_worker(i + 1, queue, browser, history, sheet, sheet_lock, stats))
        for i in range(CONCURRENCY)
    ]
    await asyncio.gather(*workers)
    log.info("Followers cycle finished: %d profiles, %d videos, %d matches",
             stats["profiles"], stats["videos"], stats["matches"])
    return stats["matches"]


# --------------------------------------------------------------------------- #
# MODE = discover
# --------------------------------------------------------------------------- #
async def collect_video_links(page: Page, scrolls: int) -> List[str]:
    found: Dict[str, None] = {}
    for _ in range(scrolls):
        hrefs = await page.eval_on_selector_all('a[href*="/video/"]', "els => els.map(e => e.getAttribute('href'))")
        for h in hrefs:
            u = clean_video_url(h or "")
            if u:
                found[u] = None
        await page.mouse.wheel(0, random.randint(1500, 2600))
        await human_pause(1.0, 2.2)
    return list(found)


async def run_discover_cycle(browser: Browser, history: History, sheet: Sheet) -> int:
    ctx = await new_context(browser)
    saved = 0
    lock = asyncio.Lock()
    try:
        urls: Dict[str, None] = {}
        sources = [f"https://www.tiktok.com/search/video?q={q}" for q in SEARCH_QUERIES]
        if USE_EXPLORE:
            sources.append("https://www.tiktok.com/explore")
        for src in sources:
            page = await ctx.new_page()
            try:
                await page.goto(src, wait_until="domcontentloaded")
                await human_pause(2, 4)
                await dismiss_popups(page)
                for u in await collect_video_links(page, DISCOVERY_SCROLLS):
                    urls[u] = None
            except Exception as e:
                log.warning("Discovery failed for %s: %s", src, e)
            finally:
                await page.close()
        links = list(urls)
        random.shuffle(links)
        todo = [u for u in links if not history.video_recently_scanned(u, RESCAN_AFTER_HOURS)]
        log.info("Discovered %d videos, %d to scan", len(links), len(todo))
        for url in todo[:MAX_VIDEOS_PER_CYCLE]:
            if STOP.is_set():
                break
            try:
                comments = await scan_video(ctx, url)
            except BlockedError:
                raise
            except Exception as e:
                log.warning("Scan failed %s: %s", url, e)
                continue
            owner = url.split("/")[3].lstrip("@") if url.count("/") >= 4 else ""
            saved += await save_matches(sheet, lock, history, comments, url, owner)
            history.mark_video(url)
            await human_pause(2, 5)
    finally:
        await ctx.close()
    return saved


# --------------------------------------------------------------------------- #
# Main loop
# --------------------------------------------------------------------------- #
async def main() -> int:
    log.info("Starting worker | mode=%s | target=@%s | concurrency=%d", MODE, TARGET_USERNAME, CONCURRENCY)
    history = History(DB_PATH)
    sheet = Sheet()
    await asyncio.to_thread(sheet.connect)

    loop = asyncio.get_running_loop()
    for sig in (signal.SIGINT, signal.SIGTERM):
        try:
            loop.add_signal_handler(sig, STOP.set)
        except NotImplementedError:
            pass

    if MAX_RUNTIME_MINUTES > 0:
        loop.call_later(MAX_RUNTIME_MINUTES * 60, STOP.set)  # graceful stop before the CI time limit
        log.info("Runtime budget: %d minutes", MAX_RUNTIME_MINUTES)

    runner = run_followers_cycle if MODE == "followers" else run_discover_cycle
    failures, cycle = 0, 0
    async with async_playwright() as pw:
        browser = await launch_browser(pw)
        try:
            while not STOP.is_set():
                cycle += 1
                try:
                    if cycle % RESTART_BROWSER_EVERY == 0:
                        log.info("Periodic browser restart")
                        await browser.close()
                        browser = await launch_browser(pw)
                        await asyncio.to_thread(history.purge, PURGE_AFTER_DAYS)

                    saved = await runner(browser, history, sheet)
                    log.info("Cycle %d done, %d new matches saved", cycle, saved)
                    failures = 0
                    if RUN_ONCE:
                        break
                    await interruptible_sleep(CYCLE_SLEEP_SECONDS + random.randint(0, 120))

                except Exception as e:
                    failures += 1
                    delay = min(BACKOFF_BASE * (2 ** (failures - 1)), BACKOFF_MAX) + random.uniform(0, 10)
                    log.error("Cycle %d failed (%s: %s). Backoff %.0fs (failure #%d)",
                              cycle, type(e).__name__, e, delay, failures)
                    try:
                        await browser.close()
                    except Exception:
                        pass
                    try:
                        browser = await launch_browser(pw)
                    except Exception as le:
                        log.error("Browser relaunch failed: %s", le)
                    await interruptible_sleep(delay)
        finally:
            try:
                await browser.close()
            except Exception:
                pass
    log.info("Worker stopped cleanly")
    return 2 if BLOCK_STATE["tripped"] else 0


if __name__ == "__main__":
    try:
        sys.exit(asyncio.run(main()) or 0)
    except KeyboardInterrupt:
        pass
